import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

def add_heading(text, level, align=WD_ALIGN_PARAGRAPH.LEFT):
    h = doc.add_heading(text, level=level)
    h.alignment = align
    return h

def add_paragraph(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(text)
    p.alignment = align
    return p

def insert_toc():
    # Skipping XML TOC to prevent Word corruption. Using placeholder text.
    p = doc.add_paragraph()
    p.add_run("<< Table of Contents will be generated here. In Word, go to References -> Table of Contents >>").bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Cover Page
doc.add_paragraph("\n\n\n\n")
add_heading("GIS-BASED WILDFIRE SUSCEPTIBILITY, RISK ZONATION AND EMERGENCY EVACUATION ACCESSIBILITY ASSESSMENT OF SANJAY GANDHI NATIONAL PARK, MUMBAI", 1, WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph("\n\n")
add_paragraph("A Project Report Submitted in Partial Fulfillment of the Requirements for the Course\nDisaster Management Mitigation", WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph("\n\n")
add_paragraph("By\n[Member 1 Name] - [Roll No]\n[Member 2 Name] - [Roll No]\n[Member 3 Name] - [Roll No]\n[Member 4 Name] - [Roll No]", WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph("\n\n")
add_paragraph("Under the Guidance of\n[Faculty Name]", WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph("\n\n")
add_paragraph("[Institute Name]\n[Semester] - [Division]", WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# DECLARATION
add_heading("DECLARATION", 2, WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("We hereby declare that this project report entitled 'GIS-Based Wildfire Susceptibility, Risk Zonation and Emergency Evacuation Accessibility Assessment of Sanjay Gandhi National Park, Mumbai' is a bona fide record of the project work carried out by us under the guidance of [Faculty Name]. We further declare that the work presented in this report is original and has not been submitted to any other institution for the award of any degree or diploma.")
add_paragraph("\n\nSignatures:\n\n___________________\n___________________\n___________________\n___________________")
doc.add_page_break()

# ACKNOWLEDGEMENT
add_heading("ACKNOWLEDGEMENT", 2, WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("We would like to express our deepest appreciation to all those who provided us with the possibility to complete this report. A special gratitude we give to our guide, [Faculty Name], whose contribution in stimulating suggestions and encouragement helped us coordinate our project especially in writing this report.")
add_paragraph("Furthermore, we would like to acknowledge with much appreciation the crucial role of the institution in providing the necessary resources. Finally, we thank our families and friends for their constant support and understanding during our project work.")
doc.add_page_break()

# ABSTRACT
add_heading("ABSTRACT", 2, WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("Wildfire in the urban-forest interface represents a complex disaster management challenge, particularly in areas like Sanjay Gandhi National Park (SGNP) in Mumbai, which is surrounded by high-density urban development. This project presents a GIS-based spatial assessment framework designed to map wildfire susceptibility, evaluate risk exposure, and prioritize emergency evacuation accessibility in and around SGNP.")
add_paragraph("The study utilizes a multi-criteria approach, combining environmental factors (such as vegetation, slope, elevation, and climate proxies) with human-related factors (such as distance to settlements and roads). These conditioning factors are synthesized into a susceptibility zonation map comprising five classes: Very Low, Low, Moderate, High, and Very High. The findings suggest that the integration of geographic information systems allows for a nuanced understanding of where wildfires are most likely to ignite and propagate. Furthermore, by overlaying susceptibility with built-up footprints, the project identifies high-risk exposure zones, enabling targeted disaster mitigation strategies.")
add_paragraph("In addition to hazard modeling, an evacuation accessibility analysis was conducted to determine potential escape routes and safe zones for populations at risk. The results demonstrate that while SGNP has critical biodiversity value, its fringes require robust emergency preparedness and structural mitigation to prevent future disasters. This study highlights the value of spatial analysis for decision-makers in developing localized, evidence-based disaster response frameworks.")
doc.add_page_break()

# TOC, LOF, LOT
add_heading("TABLE OF CONTENTS", 2, WD_ALIGN_PARAGRAPH.CENTER)
insert_toc()
add_paragraph("\n(Note: In MS Word, please go to References -> Table of Contents to insert the automatic TOC.)", WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

add_heading("LIST OF FIGURES", 2, WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("Figure 3.1: Research Methodology Workflow\nFigure 3.2: Relative Weights of Wildfire Conditioning Factors\nFigure 4.1: Study Area and Location Map\nFigure 4.2: Wildfire Susceptibility Zonation Map\nFigure 4.3: Wildfire Susceptibility Class Distribution\nFigure 4.4: Wildfire Risk and Exposure Map\nFigure 4.5: Emergency Evacuation Accessibility / Priority Map", WD_ALIGN_PARAGRAPH.LEFT)
doc.add_page_break()

add_heading("LIST OF TABLES", 2, WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("Table 2.1: Literature Review Summary\nTable 2.3: Data Sources\nTable 3.2: Analytical Framework / Illustrative Weights\nTable 5.1: Disaster Mitigation Recommendations", WD_ALIGN_PARAGRAPH.LEFT)
doc.add_page_break()

# CHAPTER 1
add_heading("CHAPTER 1: INTRODUCTION", 1)
add_heading("1.1 Background of the Study", 2)
add_paragraph("Wildfires are becoming increasingly prevalent globally and in the Indian subcontinent due to changing climatic patterns, prolonged dry spells, and intensifying human activities near forested regions. Within the context of disaster management, a wildfire is not merely an ecological event but a complex hazard capable of causing severe socio-economic disruption, loss of life, and environmental degradation. The increasing frequency of these events necessitates proactive mitigation strategies supported by spatial intelligence.")
add_paragraph("Sanjay Gandhi National Park (SGNP), located entirely within the metropolitan limits of Mumbai, presents a unique challenge. Unlike remote forest reserves, SGNP is enveloped by one of the most densely populated urban agglomerations in the world. This dense urban-forest interface increases the probability of human-induced fire ignitions while simultaneously placing a massive peripheral population at risk of fire exposure, smoke inhalation, and secondary disasters.")

add_heading("1.2 Wildfire as a Disaster Management Problem", 2)
add_paragraph("As a disaster management problem, wildfires require an understanding of both the hazard (the fire itself) and the vulnerability (the susceptibility of exposed populations and assets). In an urban-forest setting, the disaster continuum involves pre-disaster risk accumulation, the acute emergency phase of the fire, and the post-disaster recovery of the landscape. Effective management relies on spatial tools to identify which areas are most susceptible to ignition and how a fire might spread based on topography and available fuel.")
add_paragraph("Wildfires differ from other natural hazards like earthquakes because their propagation can be predicted and partially controlled if the conditioning factors are well understood. However, when a fire breaches the forest boundary and enters human settlements, it transforms into an urban conflagration, rapidly overwhelming municipal fire services. Therefore, disaster management must focus on prevention through zonation and preparedness through evacuation planning.")

add_heading("1.3 Rationale for the Study", 2)
add_paragraph("The primary rationale for this study is the pressing need to transition from reactive firefighting to proactive risk management in SGNP. Traditional methods of forest management often lack the spatial granularity required to make rapid decisions during a crisis. By employing Geographic Information Systems (GIS), this study aims to synthesize disparate environmental and anthropogenic data into actionable intelligence.")
add_paragraph("Furthermore, while broad-scale fire assessments exist for rural India, high-resolution analyses of urban national parks remain limited. The unique topography of SGNP, characterized by undulating terrain, dense vegetation, and intersecting transport corridors, demands a tailored approach. This project fills that gap by providing a conceptual spatial framework that not only identifies where fires might occur but also assesses how people can safely escape if one does.")

add_heading("1.4 Research Context", 2)
add_paragraph("The research is situated within the broader discipline of Disaster Risk Reduction (DRR) and spatial modeling. Over the past two decades, the integration of Multi-Criteria Decision Analysis (MCDA) with GIS has become the standard for hazard zonation. This study builds upon established methodologies, adapting them to the specific constraints and characteristics of the Mumbai metropolitan region.")
add_paragraph("It recognizes that risk is a product of hazard, exposure, and vulnerability. Consequently, the research context expands beyond mere susceptibility mapping to include an evaluation of built-up area exposure and the structural capacity of the road network to facilitate emergency evacuation.")

add_heading("1.5 Significance of the Study", 2)
add_paragraph("The significance of this study lies in its dual focus on environmental susceptibility and human safety. For academic purposes, it demonstrates the practical application of GIS techniques in solving real-world disaster management scenarios. For municipal and forest authorities, the conceptual maps and priority zones provide a blueprint for where to concentrate resources, clear dry brush, or establish fire lines before the dry season begins.")
add_paragraph("Additionally, the emergency evacuation accessibility component offers a novel perspective. By identifying bottlenecks and establishing priority evacuation routes, the study contributes directly to the life-safety objectives of disaster management, ensuring that vulnerable populations on the park's periphery have a planned means of egress during a catastrophic fire event.")
doc.add_page_break()

doc.save(r"d:\Aniket\majja\Project_Report_v4.docx")
print("Chapter 1 written")
